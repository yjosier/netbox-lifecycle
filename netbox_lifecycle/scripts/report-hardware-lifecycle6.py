import csv
import io
import os
from datetime import date
from django.conf import settings
from dcim.choices import DeviceStatusChoices
from dcim.models import Device
from extras.scripts import Script, ChoiceVar, MultiObjectVar
from netbox_lifecycle.models import SupportContractAssignment
from netbox_lifecycle.constants import CONTRACT_STATUS_ACTIVE

# openpyxl imports 
from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.text import RichText, Text as ChartText
from openpyxl.chart.title import Title
from openpyxl.drawing.text import (
    Paragraph, ParagraphProperties, CharacterProperties,
    RegularTextRun, Font as DrawingFont,
)
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# self.log_info(f"Device site name = {device.site.name}")
# self.log_info(f"Device site name from form = {data['site']}")
# self.log_info(f"Device status from form = {data['device_status']}")
# self.log_info(f"The device type of this device is {device.device_type}", obj=device)

# ---- Constants ----------------------------------------------------------

GLOBAL_FONT = 'Times New Roman'
BRAND_COLOR = '29ABE2'

CURRENCY_FMT = '#,##0 "CHF"'
DATE_FMT = 'dd/mm/yyyy'
LONG_DATE_FMT = 'dd mmm yyyy'
PERCENT_FMT = '0.0%'

BRAND_FILL = PatternFill('solid', start_color=BRAND_COLOR)
SECTION_HEADER_FILL = PatternFill('solid', start_color='D9E1F2')

HEADER_FONT = Font(name=GLOBAL_FONT, bold=True)
TITLE_FONT = Font(name=GLOBAL_FONT, bold=True, size=16)
SUBHEADER_FONT = Font(name=GLOBAL_FONT, bold=True, size=11)
DEFAULT_FONT = Font(name=GLOBAL_FONT)

# =========================================================================
# Data collection — the only NetBox-specific code
# =========================================================================

def collect_device_data(device_status):
    """Return a list of dicts with all the data the report needs.

    This replaces build_placeholder_data() from the standalone script.
    Output format is identical so all downstream code stays the same.
    """
    devices = Device.objects.filter(
        status=device_status,
    ).select_related('device_type', 'role', 'site')

    rows = []
    for device in devices:
        lifecycle = device.device_type.hardware_lifecycle.first()

        # renewal_price_raw = device.device_type.cf.get('Estimated_Renewal_Price')
        # # Excel needs numeric; coerce empty/None/non-numeric to 0
        # try:
        #     renewal_price = float(renewal_price_raw) if renewal_price_raw not in (None, '') else 0
        # except (TypeError, ValueError):
        #     renewal_price = 0

        active_assignments = [
            a for a in SupportContractAssignment.objects.filter(
                device=device).select_related('contract')
            if a.status == CONTRACT_STATUS_ACTIVE
        ]
        contract_id = active_assignments[0].contract.contract_id if active_assignments else ''
        contract_end = active_assignments[0].end_date if active_assignments else None

        rows.append({
            'site': device.site.name,
            'name': device.name or 'Unnamed device',
            'model': device.device_type.model,
            'role': device.role.get_root().name,
            'end_of_sale': getattr(lifecycle, 'end_of_sale', None),
            'end_of_maintenance': getattr(lifecycle, 'end_of_maintenance', None),
            'end_of_security': getattr(lifecycle, 'end_of_security', None),
            'end_of_support': getattr(lifecycle, 'end_of_support', None),
            'contract_id': contract_id,
            'contract_end': contract_end,
            'renewal_price': getattr(lifecycle, 'renewal_price', 0) or 0,
        })

    return rows


# =========================================================================
# CSV output (your original logic, adapted to use the row dicts)
# =========================================================================


def render_csv(rows):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        'Device_Site', 'Device_Name', 'Device_Type', 'Device_Role',
        'End_of_Sale', 'End_of_Maintenance', 'End_of_Security', 'End_of_Support',
        'Support_Contract', 'End_of_Contract', 'Est_Renewal_Price',
    ])

    def fmt_date(d):
        return d.strftime("%d/%m/%Y") if d else ''

    for row in rows:
        writer.writerow([
            row['site'],
            row['name'],
            row['model'],
            row['role'],
            fmt_date(row['end_of_sale']),
            fmt_date(row['end_of_maintenance']),
            fmt_date(row['end_of_security']),
            fmt_date(row['end_of_support']),
            row['contract_id'],
            fmt_date(row['contract_end']),
            # CSV preserves empty string for missing prices, unlike Excel
            row['renewal_price'] if row['renewal_price'] else '',
        ])
    return output.getvalue()


# =========================================================================
# Excel output 
# =========================================================================

# ---- Helpers ------------------------------------------------------------

def style_header_cells(ws, coords):
    """Apply section-header style (light blue fill + bold) to a list of coords."""
    for coord in coords:
        ws[coord].fill = SECTION_HEADER_FILL
        ws[coord].font = HEADER_FONT


def autosize_columns(ws, widths_by_letter):
    for letter, width in widths_by_letter.items():
        ws.column_dimensions[letter].width = width


def apply_brand_frame(ws, top_row, bottom_row, left_col, right_col):
    """Paint the brand-color frame: top + bottom rows + left + right columns."""
    for row in range(top_row, bottom_row + 1):
        ws.cell(row=row, column=left_col).fill = BRAND_FILL
        ws.cell(row=row, column=right_col).fill = BRAND_FILL
    for col in range(left_col, right_col + 1):
        ws.cell(row=top_row, column=col).fill = BRAND_FILL
        ws.cell(row=bottom_row, column=col).fill = BRAND_FILL


def apply_global_font(wb):
    """Walk every cell and apply Times New Roman, preserving bold/size/color."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                existing = cell.font
                if existing.name == GLOBAL_FONT:
                    continue
                cell.font = Font(
                    name=GLOBAL_FONT,
                    bold=existing.bold,
                    italic=existing.italic,
                    size=existing.size,
                    color=existing.color,
                    underline=existing.underline,
                )


def make_chart_title_at_bottom(text):
    """Build a chart title positioned at the bottom in Times New Roman."""
    char_props = CharacterProperties(
        latin=DrawingFont(typeface=GLOBAL_FONT),
        b=True, sz=1200,
    )
    para_props = ParagraphProperties(defRPr=char_props)
    run = RegularTextRun(rPr=char_props, t=text)
    paragraph = Paragraph(pPr=para_props, r=[run])
    rich_body = RichText(p=[paragraph])

    title = Title(overlay=False)
    title.tx = ChartText(rich=rich_body)
    title.layout = Layout(manualLayout=ManualLayout(
        xMode='edge', yMode='edge',
        x=0.25, y=0.88,
    ))
    return title

# ---- Formula building ---------------------------------------------------

def _bucket_formula(category_col, category_value, period, table_ref='raw_data'):
    """SUMPRODUCT for spend in a specific period."""
    base = (
        f'({table_ref}[{category_col}]="{category_value}")'
        f'*({table_ref}[End_of_Support]<>"")'
    )
    if period == 'overdue_and_current':
        date_filter = f'({table_ref}[End_of_Support]<DATE(YEAR(TODAY())+1,1,1))'
    else:
        offset = int(period.split('+')[1])
        date_filter = (
            f'({table_ref}[End_of_Support]>=DATE(YEAR(TODAY())+{offset},1,1))'
            f'*({table_ref}[End_of_Support]<DATE(YEAR(TODAY())+{offset+1},1,1))'
        )
    return f'=SUMPRODUCT({base}*{date_filter}*{table_ref}[Renewal_Price])'

# ---- raw_data sheet -----------------------------------------------------

def build_raw_data_sheet(wb, rows):
    ws = wb.create_sheet('raw_data')

    headers = [
        'Device_Site', 'Device_Name', 'Device_Type', 'Device_Role',
        'End of Sale', 'End of Maintenance Update', 'End of security patches',
        'End_of_Support', 'Contract support', 'End of contract support',
        'Renewal_Price',
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = SECTION_HEADER_FILL
        cell.font = HEADER_FONT

    date_cols = {5, 6, 7, 8, 10}
    price_col = 11

    for r, row in enumerate(rows, start=2):
        ws.cell(row=r, column=1, value=row['site'])
        ws.cell(row=r, column=2, value=row['name'])
        ws.cell(row=r, column=3, value=row['model'])
        ws.cell(row=r, column=4, value=row['role'])
        ws.cell(row=r, column=5, value=row['end_of_sale'])
        ws.cell(row=r, column=6, value=row['end_of_maintenance'])
        ws.cell(row=r, column=7, value=row['end_of_security'])
        ws.cell(row=r, column=8, value=row['end_of_support'])
        ws.cell(row=r, column=9, value=row['contract_id'])
        ws.cell(row=r, column=10, value=row['contract_end'])
        ws.cell(row=r, column=11, value=row['renewal_price'])

        for col in date_cols:
            cell = ws.cell(row=r, column=col)
            if cell.value is not None:
                cell.number_format = DATE_FMT
        ws.cell(row=r, column=price_col).number_format = CURRENCY_FMT

    last_row = len(rows) + 1
    last_col = get_column_letter(len(headers))
    table = Table(displayName="raw_data", ref=f"A1:{last_col}{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)

    autosize_columns(ws, {
        'A': 16, 'B': 16, 'C': 14, 'D': 14, 'E': 14, 'F': 18,
        'G': 18, 'H': 14, 'I': 14, 'J': 18, 'K': 14,
    })
    ws.freeze_panes = 'A2'
    return ws

# ---- detail_forecast sheet ---------------------------------------------

def _build_forecast_block(ws, start_row, table_name, raw_data_column,
                          display_header, category_values, today,
                          add_annual_cumulative=False):
    """Build a forecast table starting at start_row. Optionally append the
    annual + cumulative rows below it.

    Returns (table_last_row, annual_row_or_None)."""
    year = today.year
    overdue_header = f'Overdue & EoL this year ({year})'
    headers = [display_header, overdue_header,
               f'FY {year + 1}', f'FY {year + 2}', f'FY {year + 3}', 'Total']

    for c, h in enumerate(headers, start=1):
        ws.cell(row=start_row, column=c, value=h)

    periods = ['overdue_and_current', 'fy+1', 'fy+2', 'fy+3']
    body_start = start_row + 1
    for r_offset, value in enumerate(category_values):
        r = body_start + r_offset
        ws.cell(row=r, column=1, value=value)
        for col_idx, period in enumerate(periods, start=2):
            cell = ws.cell(row=r, column=col_idx,
                           value=_bucket_formula(raw_data_column, value, period))
            cell.number_format = CURRENCY_FMT
        last_period_header = f'FY {year + 3}'
        total_cell = ws.cell(
            row=r, column=6,
            value=(f'=SUM({table_name}[[#This Row],'
                   f'[{overdue_header}]:[{last_period_header}]])')
        )
        total_cell.number_format = CURRENCY_FMT

    table_last_row = body_start + len(category_values) - 1
    last_col = get_column_letter(len(headers))
    table = Table(
        displayName=table_name,
        ref=f"A{start_row}:{last_col}{table_last_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)

    annual_row = None
    if add_annual_cumulative:
        annual_row = table_last_row + 2
        cumulative_row = annual_row + 1

        ws.cell(row=annual_row, column=1, value='Annual spend').font = SUBHEADER_FONT
        ws.cell(row=cumulative_row, column=1,
                value='Cumulative spend').font = SUBHEADER_FONT

        for col_idx in range(2, 6):
            col_letter = get_column_letter(col_idx)
            annual_cell = ws.cell(
                row=annual_row, column=col_idx,
                value=f'=SUM({col_letter}{body_start}:{col_letter}{table_last_row})'
            )
            annual_cell.number_format = CURRENCY_FMT
            annual_cell.font = SUBHEADER_FONT

            if col_idx == 2:
                cum_formula = f'={col_letter}{annual_row}'
            else:
                prev = get_column_letter(col_idx - 1)
                cum_formula = f'={prev}{cumulative_row}+{col_letter}{annual_row}'
            cum_cell = ws.cell(row=cumulative_row, column=col_idx, value=cum_formula)
            cum_cell.number_format = CURRENCY_FMT
            cum_cell.font = SUBHEADER_FONT

    return table_last_row, annual_row


def build_detail_forecast_sheet(wb, today, unique_sites, unique_roles):
    ws = wb.create_sheet('detail_forecast')

    # Per-site section title (merged A1:F1)
    ws['A1'] = 'Per-site forecast'
    ws['A1'].font = SUBHEADER_FONT
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A1:F1')

    sites_last_row, _ = _build_forecast_block(
        ws, start_row=2,
        table_name='by_site_forecast',
        raw_data_column='Device_Site',
        display_header='Device_Site',
        category_values=unique_sites,
        today=today,
        add_annual_cumulative=False,
    )

    # Per device-role section title (merged) — placed two rows below the site table
    role_section_row = sites_last_row + 3
    ws.cell(row=role_section_row, column=1, value='Per device-role forecast').font = SUBHEADER_FONT
    ws.cell(row=role_section_row, column=1).alignment = Alignment(horizontal='center')
    ws.merge_cells(start_row=role_section_row, start_column=1,
                   end_row=role_section_row, end_column=6)

    role_last_row, role_annual_row = _build_forecast_block(
        ws, start_row=role_section_row + 1,
        table_name='by_role_forecast',
        raw_data_column='Device_Role',
        display_header='Device_Role',
        category_values=unique_roles,
        today=today,
        add_annual_cumulative=True,
    )

    autosize_columns(ws, {'A': 22, 'B': 30, 'C': 14, 'D': 14, 'E': 14, 'F': 14})
    return ws, role_annual_row

# ---- Exec summary sheet -------------------------------------------------

def build_exec_summary(wb, today, role_annual_row_in_detail):
    ws = wb.create_sheet('Exec summary', 0)
    ws.sheet_view.showGridLines = False
    year = today.year

    # ---- Title block (B2 merged across B2:N2) ----
    ws['B2'] = 'Hardware Lifecycle — Financial Forecast'
    ws['B2'].font = TITLE_FONT
    ws['B2'].alignment = Alignment(horizontal='center', vertical='center')
    ws.merge_cells('B2:N2')

    ws['B3'] = 'Generated:'
    ws['C3'] = '=TODAY()'
    ws['C3'].number_format = LONG_DATE_FMT

    # ---- Headline metrics block ----
    ws['B6'] = 'Headline metrics'
    style_header_cells(ws, ['B6', 'C6'])

    ws['B7'] = 'Total devices audited'
    ws['C7'] = '=COUNTA(raw_data[Device_Name])'

    ws['B8'] = 'Devices without an EoL date'
    ws['C8'] = '=ROWS(raw_data[End_of_Support])-COUNT(raw_data[End_of_Support])'

    ws['B9'] = 'Blind spot (% missing EoL)'
    ws['C9'] = '=C8/C7'
    ws['C9'].number_format = PERCENT_FMT

    ws['B10'] = 'Devices past EoL today'
    ws['C10'] = '=COUNTIFS(raw_data[End_of_Support],"<"&TODAY(),raw_data[End_of_Support],"<>")'

    # ---- Spend forecast block ----
    ws['B13'] = 'Spend forecast'
    style_header_cells(ws, ['B13', 'C13'])

    ws['B14'] = f'Estimated spend — overdue & current FY ({year})'
    ws['C14'] = f"=detail_forecast!B{role_annual_row_in_detail}"
    ws['C14'].number_format = CURRENCY_FMT

    spend_rows = [
        (15, f'Estimated spend — FY ({year + 1})', 'C'),
        (16, f'Estimated spend — FY ({year + 2})', 'D'),
        (17, f'Estimated spend — FY ({year + 3})', 'E'),
    ]
    for row, label, col_letter in spend_rows:
        ws.cell(row=row, column=2, value=label)
        cell = ws.cell(row=row, column=3,
                       value=f"=detail_forecast!{col_letter}{role_annual_row_in_detail}")
        cell.number_format = CURRENCY_FMT

    # ---- Top-3 sites block ----
    ws['F5'] = 'Top spending sites (3-year window)'
    ws['F5'].font = SUBHEADER_FONT

    ws['F6'], ws['G6'], ws['H6'] = 'Rank', 'Site', 'Total spend'
    style_header_cells(ws, ['F6', 'G6', 'H6'])

    for rank in range(1, 4):
        r = 6 + rank
        ws.cell(row=r, column=6, value=rank)
        ws.cell(row=r, column=7, value=(
            f'=IFERROR(INDEX(by_site_forecast[Device_Site],'
            f'MATCH(LARGE(by_site_forecast[Total],{rank}),by_site_forecast[Total],0)),"-")'
        ))
        ws.cell(row=r, column=8, value=(
            f'=IFERROR(LARGE(by_site_forecast[Total],{rank}),0)'
        ))
        ws.cell(row=r, column=8).number_format = CURRENCY_FMT

    ws['F10'] = '—'
    ws['G10'] = 'Others'
    ws['H10'] = (
        '=SUM(by_site_forecast[Total])'
        '-IFERROR(LARGE(by_site_forecast[Total],1),0)'
        '-IFERROR(LARGE(by_site_forecast[Total],2),0)'
        '-IFERROR(LARGE(by_site_forecast[Total],3),0)'
    )
    ws['H10'].number_format = CURRENCY_FMT

    # ---- Top-3 categories block ----
    ws['K5'] = 'Top spending categories (3-year window)'
    ws['K5'].font = SUBHEADER_FONT

    ws['K6'], ws['L6'], ws['M6'] = 'Rank', 'Category', 'Total spend'
    style_header_cells(ws, ['K6', 'L6', 'M6'])

    for rank in range(1, 4):
        r = 6 + rank
        ws.cell(row=r, column=11, value=rank)
        ws.cell(row=r, column=12, value=(
            f'=IFERROR(INDEX(by_role_forecast[Device_Role],'
            f'MATCH(LARGE(by_role_forecast[Total],{rank}),by_role_forecast[Total],0)),"-")'
        ))
        ws.cell(row=r, column=13, value=(
            f'=IFERROR(LARGE(by_role_forecast[Total],{rank}),0)'
        ))
        ws.cell(row=r, column=13).number_format = CURRENCY_FMT

    ws['K10'] = '—'
    ws['L10'] = 'Others'
    ws['M10'] = (
        '=SUM(by_role_forecast[Total])'
        '-IFERROR(LARGE(by_role_forecast[Total],1),0)'
        '-IFERROR(LARGE(by_role_forecast[Total],2),0)'
        '-IFERROR(LARGE(by_role_forecast[Total],3),0)'
    )
    ws['M10'].number_format = CURRENCY_FMT

    # ---- Pie chart: sites (anchored at E13) ----
    pie_sites = PieChart()
    pie_sites.title = make_chart_title_at_bottom('Top spending sites — 3-year window')
    labels = Reference(ws, min_col=7, min_row=7, max_row=10)
    data = Reference(ws, min_col=8, min_row=7, max_row=10)
    pie_sites.add_data(data, titles_from_data=False)
    pie_sites.set_categories(labels)
    pie_sites.dataLabels = DataLabelList(
        showCatName=True, showVal=True, showPercent=True, showSerName=False,
    )
    pie_sites.legend = None
    pie_sites.height = 8
    pie_sites.width = 12
    ws.add_chart(pie_sites, 'E13')

    # ---- Pie chart: categories (anchored at J13) ----
    pie_roles = PieChart()
    pie_roles.title = make_chart_title_at_bottom('Top spending categories — 3-year window')
    labels = Reference(ws, min_col=12, min_row=7, max_row=10)
    data = Reference(ws, min_col=13, min_row=7, max_row=10)
    pie_roles.add_data(data, titles_from_data=False)
    pie_roles.set_categories(labels)
    pie_roles.dataLabels = DataLabelList(
        showCatName=True, showVal=True, showPercent=True, showSerName=False,
    )
    pie_roles.legend = None
    pie_roles.height = 8
    pie_roles.width = 12
    ws.add_chart(pie_roles, 'J13')

    # ---- Definitions and disclaimer ----
    ws['B33'] = 'Definitions & disclaimer'
    ws['B33'].font = SUBHEADER_FONT

    disclaimer_text = (
        "End of Life (EoL): For the purposes of this report, EoL refers to the "
        "manufacturer's declared end of support date — the last day on which "
        "the vendor will provide official assistance for the hardware. Other "
        "vendor milestones such as end of sale, end of maintenance, and end of "
        "security patches are tracked separately in the raw_data tab and may "
        "occur earlier.\n"
        "Disclaimer: This forecast assumes a one-for-one replacement of every "
        "device reaching end of support within the indicated fiscal year, "
        "valued at its current renewal price. Actual capital expenditure may "
        "differ materially due to extended-support agreements, planned "
        "decommissioning, consolidation initiatives, vendor pricing changes, "
        "and procurement timing. The figures presented are intended for "
        "planning and budgetary discussion, not as a commitment to spend."
    )
    ws['B34'] = disclaimer_text
    ws['B34'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells('B34:N34')
    ws.row_dimensions[34].height = 65
    # ---- Brand-color frame (rows 1-35, columns A-O) ----
    apply_brand_frame(ws, top_row=1, bottom_row=35, left_col=1, right_col=15)

    autosize_columns(ws, {
        'A': 2, 'B': 44, 'C': 16, 'D': 10, 'E': 10, 'F': 10,
        'G': 18, 'H': 16, 'I': 15, 'J': 15, 'K': 10, 'L': 18,
        'M': 16, 'N': 10, 'O': 2,
    })

    return ws


def generate_report(rows, output_path, today=None):
    today = today or date.today()

    wb = Workbook()
    wb.remove(wb.active)

    unique_sites = sorted({r['site'] for r in rows})
    unique_roles = sorted({r['role'] for r in rows})

    build_raw_data_sheet(wb, rows)
    detail_ws, role_annual_row = build_detail_forecast_sheet(
        wb, today, unique_sites, unique_roles
    )
    build_exec_summary(wb, today, role_annual_row)

    desired_order = ['Exec summary', 'detail_forecast', 'raw_data']
    for idx, name in enumerate(desired_order):
        current_idx = wb.sheetnames.index(name)
        if current_idx != idx:
            wb.move_sheet(name, offset=idx - current_idx)

    apply_global_font(wb)

    wb.save(output_path)
    return output_path


# =========================================================================
# The Custom Script class
# =========================================================================


class HardwareLifecycleReport(Script):
    class Meta:
        name = "Hardware Lifecycle Report"
        description = "Generate a financial forecast of hardware lifecycle for budget planning"

    device_status = ChoiceVar(
        choices=DeviceStatusChoices,
        default=DeviceStatusChoices.STATUS_ACTIVE,
        label="Device status",
        description="Only include devices with this status",
    )

    output_format = ChoiceVar(
        choices=(
            ('xlsx', 'Excel (xlsx) — full styled report'),
            ('csv', 'CSV — raw data only'),
        ),
        default='xlsx',
        label="Output format",
    )

    def run(self, data, commit):
        rows = collect_device_data(data['device_status'])

        if not rows:
            self.log_warning("No devices matched. Nothing to export.")
            return "No data."

        self.log_info(f"Collected {len(rows)} device(s).")

        # Determine output filename and path
        timestamp = date.today().isoformat()
        ext = data['output_format']
        filename = f"lifecycle_report_{timestamp}.{ext}"

        output_dir = os.path.join(settings.MEDIA_ROOT, 'lifecycle_reports')
        os.makedirs(output_dir, exist_ok=True)
        filepath = os.path.join(output_dir, filename)

        if ext == 'csv':
            content = render_csv(rows)
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(content)
        else:
            # Excel: generate_report() writes the file directly
            generate_report(rows, filepath)

        # Build the user-facing download URL
        media_url = settings.MEDIA_URL.rstrip('/')
        download_url = f"{media_url}/lifecycle_reports/{filename}"
        self.log_success(f"Report saved. Download: <a href='{download_url}'>{filename}</a>")

        return f"Wrote {filepath}"